from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "电商系统功能演示.pptx"

FONT = "Microsoft YaHei"
BG = RGBColor(247, 248, 250)
INK = RGBColor(24, 34, 48)
MUTED = RGBColor(102, 112, 133)
TEAL = RGBColor(15, 118, 110)
BLUE = RGBColor(37, 99, 235)
PANEL = RGBColor(255, 255, 255)
LINE = RGBColor(223, 227, 232)


def main() -> None:
    prs = Presentation()
    prs.slide_width = Cm(33.867)
    prs.slide_height = Cm(19.05)

    cover(prs)
    storefront(prs)
    admin(prs)
    order_flow(prs)
    operations(prs)
    demo_path(prs)

    prs.save(OUTPUT)
    print(f"generated: {OUTPUT}")
    print(f"slides: {len(prs.slides)}")


def slide(prs: Presentation, title: str, subtitle: str = ""):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    text(s, title, Cm(1.5), Cm(0.9), Cm(25), Cm(0.9), 25, bold=True)
    line(s, Cm(1.5), Cm(2.05), Cm(30.8))
    if subtitle:
        text(s, subtitle, Cm(1.55), Cm(2.32), Cm(29), Cm(0.55), 11, MUTED)
    return s


def cover(prs: Presentation) -> None:
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = RGBColor(241, 245, 249)
    text(s, "数据驱动电商系统", Cm(1.7), Cm(2.4), Cm(22), Cm(1.1), 32, bold=True)
    text(s, "项目功能演示", Cm(1.75), Cm(3.85), Cm(10), Cm(0.8), 18, TEAL, bold=True)
    text(
        s,
        "一个可运行的中文前台商城 + 中文运营后台，覆盖商品、购物车、订单、支付、物流、售后和报表。",
        Cm(1.75),
        Cm(5.0),
        Cm(25),
        Cm(0.75),
        13,
        MUTED,
    )
    cards = [
        ("前台商城", "浏览商品、加购、下单"),
        ("运营后台", "商品、订单、客户、售后"),
        ("交易闭环", "支付、发货、退款"),
        ("数据运营", "报表、审计、监控"),
    ]
    for i, (title, body) in enumerate(cards):
        card(s, Cm(1.75 + i * 7.75), Cm(8.1), Cm(6.8), Cm(3.0), title, body)


def storefront(prs: Presentation) -> None:
    s = slide(prs, "前台商城功能", "给买家使用，完成从浏览商品到提交订单的流程。")
    items = [
        ("商品浏览", "展示已上架商品、价格、库存、图片和描述。"),
        ("买家身份", "通过邮箱创建或载入买家，方便测试不同客户。"),
        ("购物车", "使用 Redis 保存购物车，支持选择商品和数量。"),
        ("提交订单", "购物车结算后生成订单，并展示订单确认页。"),
    ]
    feature_grid(s, items)


def admin(prs: Presentation) -> None:
    s = slide(prs, "中文运营后台功能", "给运营人员使用，管理商品、订单、客户和售后。")
    items = [
        ("商品管理", "新增商品、编辑价格库存、上下架、上传商品图片。"),
        ("库存管理", "设置低库存阈值，记录每次库存调整原因。"),
        ("客户管理", "查看客户列表、搜索客户、查看历史订单和累计消费。"),
        ("订单管理", "按状态、客户、日期筛选订单，支持 CSV 导出。"),
    ]
    feature_grid(s, items)


def order_flow(prs: Presentation) -> None:
    s = slide(prs, "订单与交易流程", "演示重点是完整交易闭环，而不是单个页面。")
    steps = ["商品上架", "买家加购", "提交订单", "发起支付", "后台发货", "售后退款"]
    y = Cm(6.8)
    for i, step in enumerate(steps):
        x = Cm(1.7 + i * 5.1)
        node(s, step, x, y, Cm(4.0), Cm(1.55), TEAL if i < 3 else BLUE)
        if i < len(steps) - 1:
            arrow(s, x + Cm(4.0), y + Cm(0.78), x + Cm(5.0), y + Cm(0.78))
    text(s, "后台可模拟支付成功，也预留真实支付 checkout 和签名 Webhook。", Cm(2.0), Cm(10.4), Cm(29), Cm(0.7), 13, MUTED)
    text(s, "发货后记录承运商、物流单号和追踪链接；售后完成后可标记退款。", Cm(2.0), Cm(11.45), Cm(29), Cm(0.7), 13, MUTED)


def operations(prs: Presentation) -> None:
    s = slide(prs, "运营与安全能力", "除了业务流程，还提供后台管理和运维基础能力。")
    items = [
        ("权限控制", "管理员登录，Owner / Operator 角色权限矩阵。"),
        ("API Key", "外部系统通过 API Key 调用接口，可启停和限流。"),
        ("审计日志", "记录登录失败、商品、库存、订单、支付、售后等关键操作。"),
        ("监控状态", "提供健康检查、运行状态和 Prometheus 风格指标。"),
    ]
    feature_grid(s, items)


def demo_path(prs: Presentation) -> None:
    s = slide(prs, "推荐演示顺序", "按这个顺序讲，5 分钟内能把项目能力讲清楚。")
    steps = [
        "打开 /store，展示前台商品列表和买家购物车。",
        "进入 /admin，展示中文运营后台首页。",
        "创建或编辑商品，上传商品图片。",
        "前台加入购物车并提交订单。",
        "后台查看订单，模拟支付、填写物流、处理售后。",
        "展示报表中心、审计日志、API Key 和监控端点。",
    ]
    for i, item in enumerate(steps, start=1):
        y = Cm(3.25 + (i - 1) * 2.0)
        circle(s, str(i), Cm(1.9), y)
        text(s, item, Cm(3.25), y + Cm(0.04), Cm(27), Cm(0.55), 14, INK)


def feature_grid(s, items: list[tuple[str, str]]) -> None:
    for i, (title, body) in enumerate(items):
        x = Cm(1.7 + (i % 2) * 15.9)
        y = Cm(3.45 + (i // 2) * 5.2)
        card(s, x, y, Cm(14.4), Cm(3.9), title, body)


def card(s, x, y, w, h, title, body) -> None:
    shape = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = PANEL
    shape.line.color.rgb = LINE
    text(s, title, x + Cm(0.45), y + Cm(0.45), w - Cm(0.9), Cm(0.6), 15, TEAL, bold=True)
    text(s, body, x + Cm(0.45), y + Cm(1.45), w - Cm(0.9), Cm(1.1), 12, MUTED)


def node(s, label, x, y, w, h, color) -> None:
    shape = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.color.rgb = color
    tf = shape.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = label
    run.font.name = FONT
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = RGBColor(255, 255, 255)


def circle(s, label, x, y) -> None:
    shape = s.shapes.add_shape(MSO_SHAPE.OVAL, x, y, Cm(0.86), Cm(0.86))
    shape.fill.solid()
    shape.fill.fore_color.rgb = TEAL
    shape.line.color.rgb = TEAL
    text(s, label, x, y + Cm(0.11), Cm(0.86), Cm(0.4), 10, RGBColor(255, 255, 255), bold=True, align=PP_ALIGN.CENTER)


def arrow(s, x1, y1, x2, y2) -> None:
    connector = s.shapes.add_connector(1, x1, y1, x2, y2)
    connector.line.color.rgb = MUTED
    connector.line.width = Pt(1.3)


def line(s, x, y, w) -> None:
    shape = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, Cm(0.03))
    shape.fill.solid()
    shape.fill.fore_color.rgb = TEAL
    shape.line.color.rgb = TEAL


def text(s, value, x, y, w, h, size, color=INK, bold=False, align=PP_ALIGN.LEFT):
    box = s.shapes.add_textbox(x, y, w, h)
    frame = box.text_frame
    frame.clear()
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    p = frame.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = value
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


if __name__ == "__main__":
    main()
