from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "电商系统项目演示.pptx"

FONT = "Microsoft YaHei"
INK = RGBColor(24, 34, 48)
MUTED = RGBColor(102, 112, 133)
TEAL = RGBColor(15, 118, 110)
TEAL_DARK = RGBColor(17, 94, 89)
BG = RGBColor(247, 248, 250)
PANEL = RGBColor(255, 255, 255)
LINE = RGBColor(223, 227, 232)
BLUE = RGBColor(37, 99, 235)
AMBER = RGBColor(180, 83, 9)


def main() -> None:
    prs = Presentation()
    prs.slide_width = Cm(33.867)
    prs.slide_height = Cm(19.05)

    add_cover(prs)
    add_positioning(prs)
    add_stack(prs)
    add_capabilities(prs)
    add_architecture(prs)
    add_flow(prs)
    add_production(prs)
    add_quality(prs)
    add_demo(prs)
    add_summary(prs)

    prs.save(OUTPUT)
    print(f"generated: {OUTPUT}")
    print(f"slides: {len(prs.slides)}")


def blank(prs: Presentation, title: str, subtitle: str = ""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = BG
    add_text(slide, title, Cm(1.4), Cm(0.8), Cm(18), Cm(1.1), 25, bold=True)
    add_line(slide, Cm(1.4), Cm(2.0), Cm(31.0), TEAL)
    if subtitle:
        add_text(slide, subtitle, Cm(1.45), Cm(2.25), Cm(25), Cm(0.7), 11, MUTED)
    add_text(slide, "数据驱动电商系统平台", Cm(25.2), Cm(17.9), Cm(7), Cm(0.45), 8, MUTED)
    return slide


def add_cover(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(241, 245, 249)
    add_text(slide, "数据驱动电商系统平台", Cm(1.6), Cm(2.1), Cm(23), Cm(1.2), 31, bold=True)
    add_text(slide, "Python FastAPI + PostgreSQL + Redis 全栈电商演示", Cm(1.7), Cm(3.55), Cm(24), Cm(0.8), 16, TEAL_DARK)
    add_text(
        slide,
        "覆盖前台商城、中文运营后台、商品库存、订单支付、物流售后、权限审计、监控与 Windows 原生部署。",
        Cm(1.7),
        Cm(4.65),
        Cm(24),
        Cm(0.8),
        12,
        MUTED,
    )
    items = [
        ("后端框架", "FastAPI / SQLModel"),
        ("数据存储", "PostgreSQL / Redis"),
        ("页面能力", "Jinja2 中文前后台"),
        ("工程质量", "Alembic / pytest / ruff / mypy"),
    ]
    for index, (label, value) in enumerate(items):
        x = Cm(1.7 + index * 7.7)
        add_card(slide, x, Cm(7.2), Cm(6.8), Cm(3.1), label, value, TEAL if index == 0 else BLUE)
    add_text(slide, "独立全栈开发", Cm(1.7), Cm(13.7), Cm(8), Cm(0.7), 13, bold=True)
    add_text(slide, "本地可运行、可演示，并预留真实支付/物流接入与生产部署能力。", Cm(1.7), Cm(14.5), Cm(18), Cm(0.6), 11, MUTED)


def add_positioning(prs: Presentation) -> None:
    slide = blank(prs, "项目定位", "围绕电商核心业务，构建可演示、可扩展、可生产化的全栈系统。")
    bullets = [
        ("解决什么问题", "统一管理商品、库存、客户、订单、支付、售后等电商核心数据。"),
        ("面向哪些角色", "运营人员使用后台处理商品、订单、客户、售后；买家通过前台完成浏览和下单。"),
        ("做到什么程度", "已形成前台商城 + 后台运营 + API 接入 + 监控部署的完整项目骨架。"),
        ("技术侧价值", "用 FastAPI、SQLModel、PostgreSQL、Redis 打通业务闭环，并配套迁移、测试和类型检查。"),
    ]
    add_bullet_panel(slide, Cm(1.6), Cm(3.2), Cm(30.4), Cm(11.8), bullets)


def add_stack(prs: Presentation) -> None:
    slide = blank(prs, "技术栈", "采用 Python 全栈黄金组合，兼顾开发效率、数据建模和本地部署可行性。")
    stacks = [
        ("后端 API", ["FastAPI", "Pydantic", "依赖注入", "API Key 鉴权"]),
        ("数据层", ["SQLModel", "PostgreSQL", "Alembic 迁移", "Repository 分层"]),
        ("缓存与稳定性", ["Redis 购物车", "API 限流", "访问日志", "健康检查"]),
        ("页面与测试", ["Jinja2 中文页面", "pytest", "ruff", "mypy", "Playwright"]),
    ]
    for i, (title, tags) in enumerate(stacks):
        x = Cm(1.5 + (i % 2) * 15.8)
        y = Cm(3.1 + (i // 2) * 5.6)
        add_tag_card(slide, x, y, Cm(14.5), Cm(4.4), title, tags)


def add_capabilities(prs: Presentation) -> None:
    slide = blank(prs, "核心功能", "覆盖电商从商品上架到订单履约、售后退款的主要流程。")
    cols = [
        ("前台商城", ["商品浏览", "买家创建/载入", "Redis 购物车", "提交订单", "订单确认"]),
        ("运营后台", ["商品管理", "库存调整", "订单筛选", "客户管理", "报表中心"]),
        ("交易履约", ["支付会话", "Webhook 回调", "发货物流", "售后申请", "退款状态"]),
        ("安全治理", ["管理员登录", "角色权限", "API Key", "接口限流", "审计日志"]),
    ]
    for i, (title, items) in enumerate(cols):
        add_list_card(slide, Cm(1.4 + i * 8.0), Cm(3.0), Cm(7.2), Cm(11.3), title, items)


def add_architecture(prs: Presentation) -> None:
    slide = blank(prs, "系统架构", "按前台、后台、API、服务层、数据层拆分，便于后续扩展真实支付和物流。")
    nodes = [
        ("前台商城\n/store", 1.6, 3.3, TEAL),
        ("运营后台\n/admin", 1.6, 7.2, BLUE),
        ("开放 API\n/api/v1", 1.6, 11.1, AMBER),
        ("FastAPI 路由层", 10.7, 5.0, TEAL_DARK),
        ("业务服务层\nCart / Orders / Payments / Shipping", 18.1, 5.0, BLUE),
        ("PostgreSQL\n业务数据", 26.1, 3.4, TEAL),
        ("Redis\n购物车/限流", 26.1, 8.0, AMBER),
        ("Ops 监控\n/ops/metrics", 26.1, 12.6, BLUE),
    ]
    positions: dict[str, tuple[float, float]] = {}
    for label, x, y, color in nodes:
        positions[label] = (x, y)
        add_node(slide, label, Cm(x), Cm(y), Cm(6.2), Cm(2.0), color)
    for source in nodes[:3]:
        add_arrow(slide, Cm(source[1] + 6.2), Cm(source[2] + 1.0), Cm(10.7), Cm(6.0))
    add_arrow(slide, Cm(16.9), Cm(6.0), Cm(18.1), Cm(6.0))
    add_arrow(slide, Cm(24.3), Cm(5.8), Cm(26.1), Cm(4.4))
    add_arrow(slide, Cm(24.3), Cm(6.3), Cm(26.1), Cm(9.0))
    add_arrow(slide, Cm(24.3), Cm(6.8), Cm(26.1), Cm(13.6))


def add_flow(prs: Presentation) -> None:
    slide = blank(prs, "核心业务流程", "买家侧和运营侧共用同一套数据模型，形成完整闭环。")
    steps = [
        "商品上架",
        "前台浏览",
        "加入购物车",
        "提交订单",
        "支付回调",
        "发货物流",
        "售后退款",
        "报表审计",
    ]
    y = Cm(6.2)
    for i, step in enumerate(steps):
        x = Cm(1.2 + i * 4.0)
        add_node(slide, step, x, y, Cm(3.2), Cm(1.35), TEAL if i < 4 else BLUE)
        if i < len(steps) - 1:
            add_arrow(slide, x + Cm(3.2), y + Cm(0.68), x + Cm(4.0), y + Cm(0.68))
    add_text(slide, "后台能力：商品管理、库存调整、订单处理、客户管理、售后处理、报表中心", Cm(1.5), Cm(10.0), Cm(29), Cm(0.7), 12, MUTED)
    add_text(slide, "前台能力：中文商城、买家身份、Redis 购物车、订单确认、支付入口", Cm(1.5), Cm(11.2), Cm(29), Cm(0.7), 12, MUTED)


def add_production(prs: Presentation) -> None:
    slide = blank(prs, "生产化能力", "不依赖 Docker，兼容 Windows 11 家庭版的本地服务部署方式。")
    bullets = [
        ("支付接入预留", "支付 checkout 会话、签名 Webhook、幂等事件表，后续可接支付宝/微信/Stripe。"),
        ("物流接入预留", "Shipment 运单模型沉淀承运商、物流单号、追踪链接和发货状态。"),
        ("监控指标", "提供 /health/ready、/ops/status、/ops/metrics，便于接入监控系统。"),
        ("部署脚本", "提供 Windows 原生 start_prod.ps1、check_prod.ps1 与生产配置样例。"),
    ]
    add_bullet_panel(slide, Cm(1.6), Cm(3.1), Cm(30.4), Cm(11.8), bullets)


def add_quality(prs: Presentation) -> None:
    slide = blank(prs, "工程质量", "项目已配套迁移、测试、类型检查和浏览器回归入口。")
    metrics = [
        ("22", "pytest 用例通过"),
        ("44", "mypy 检查源码文件"),
        ("100%", "ruff 当前通过"),
        ("9", "Alembic 数据库迁移"),
    ]
    for i, (value, label) in enumerate(metrics):
        add_metric(slide, Cm(1.7 + i * 7.8), Cm(3.4), Cm(6.6), Cm(3.4), value, label)
    bullets = [
        ("数据库演进", "所有结构变化通过 Alembic 管理，支持本地和生产环境一致迁移。"),
        ("测试覆盖", "覆盖后台页面、API 鉴权、购物车下单、前台商城、支付 Webhook、监控端点。"),
        ("代码规范", "ruff 做风格检查，mypy 做类型检查，降低后续迭代风险。"),
    ]
    add_bullet_panel(slide, Cm(1.7), Cm(8.2), Cm(30.0), Cm(6.6), bullets)


def add_demo(prs: Presentation) -> None:
    slide = blank(prs, "演示路径", "可按以下顺序进行 5 分钟项目演示。")
    items = [
        ("1", "打开 /store，展示中文前台商城、商品列表和购物车。"),
        ("2", "进入 /admin，展示中文后台的商品、订单、客户、报表和权限。"),
        ("3", "创建商品并上传图片，观察前台商品同步展示。"),
        ("4", "用前台创建买家、加入购物车、提交订单。"),
        ("5", "后台处理订单支付、发货物流、售后退款并查看审计日志。"),
        ("6", "打开 /ops/status 和 /ops/metrics 展示监控能力。"),
    ]
    for i, (num, text) in enumerate(items):
        y = Cm(3.1 + i * 2.1)
        add_circle_label(slide, num, Cm(1.8), y)
        add_text(slide, text, Cm(3.2), y - Cm(0.05), Cm(27), Cm(0.6), 13, INK)


def add_summary(prs: Presentation) -> None:
    slide = blank(prs, "项目总结", "一个可运行、可演示、可继续扩展上线的全栈电商系统。")
    add_text(
        slide,
        "项目已经完成“前台商城 + 中文运营后台 + 数据存储 + 缓存限流 + 支付物流扩展 + 监控部署”的核心闭环。",
        Cm(1.7),
        Cm(3.2),
        Cm(28.5),
        Cm(1.0),
        18,
        bold=True,
    )
    bullets = [
        ("业务完整", "商品、库存、客户、订单、支付、物流、售后、报表均有落地能力。"),
        ("技术完整", "FastAPI、SQLModel、PostgreSQL、Redis、Alembic、Jinja2、pytest 串成完整工程。"),
        ("扩展明确", "真实支付/物流只需补商户密钥和供应商 SDK，即可沿现有适配器接入。"),
    ]
    add_bullet_panel(slide, Cm(1.7), Cm(6.1), Cm(30.0), Cm(7.0), bullets)


def add_card(slide, x, y, w, h, label, value, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = PANEL
    shape.line.color.rgb = LINE
    add_text(slide, label, x + Cm(0.35), y + Cm(0.35), w - Cm(0.7), Cm(0.5), 10, MUTED)
    add_text(slide, value, x + Cm(0.35), y + Cm(1.25), w - Cm(0.7), Cm(1.0), 14, color, bold=True)


def add_tag_card(slide, x, y, w, h, title, tags):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = PANEL
    shape.line.color.rgb = LINE
    add_text(slide, title, x + Cm(0.45), y + Cm(0.35), w - Cm(0.9), Cm(0.6), 15, bold=True)
    tx, ty = x + Cm(0.45), y + Cm(1.35)
    for tag in tags:
        add_pill(slide, tag, tx, ty)
        tx += Cm(max(2.5, len(tag) * 0.34 + 1.0))
        if tx > x + w - Cm(3.0):
            tx = x + Cm(0.45)
            ty += Cm(0.85)


def add_list_card(slide, x, y, w, h, title, items):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = PANEL
    shape.line.color.rgb = LINE
    add_text(slide, title, x + Cm(0.45), y + Cm(0.45), w - Cm(0.9), Cm(0.7), 15, TEAL_DARK, bold=True)
    for i, item in enumerate(items):
        add_text(slide, f"▪ {item}", x + Cm(0.55), y + Cm(1.55 + i * 1.0), w - Cm(1.1), Cm(0.5), 11, INK)


def add_bullet_panel(slide, x, y, w, h, bullets):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = PANEL
    shape.line.color.rgb = LINE
    for i, (head, body) in enumerate(bullets):
        yy = y + Cm(0.65 + i * 2.45)
        add_text(slide, f"▪ {head}", x + Cm(0.6), yy, w - Cm(1.2), Cm(0.55), 14, INK, bold=True)
        add_text(slide, body, x + Cm(1.05), yy + Cm(0.75), w - Cm(1.55), Cm(0.7), 12, MUTED)


def add_metric(slide, x, y, w, h, value, label):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = PANEL
    shape.line.color.rgb = LINE
    add_text(slide, value, x, y + Cm(0.45), w, Cm(1.2), 30, TEAL, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, label, x, y + Cm(2.05), w, Cm(0.6), 11, MUTED, align=PP_ALIGN.CENTER)


def add_node(slide, text, x, y, w, h, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.color.rgb = color
    box = shape.text_frame
    box.clear()
    p = box.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.color.rgb = RGBColor(255, 255, 255)


def add_pill(slide, text, x, y):
    w = Cm(max(2.2, len(text) * 0.32 + 0.8))
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, Cm(0.6))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(236, 253, 245)
    shape.line.color.rgb = RGBColor(167, 243, 208)
    add_text(slide, text, x, y + Cm(0.08), w, Cm(0.4), 8, TEAL_DARK, bold=True, align=PP_ALIGN.CENTER)


def add_circle_label(slide, text, x, y):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y, Cm(0.85), Cm(0.85))
    shape.fill.solid()
    shape.fill.fore_color.rgb = TEAL
    shape.line.color.rgb = TEAL
    add_text(slide, text, x, y + Cm(0.12), Cm(0.85), Cm(0.4), 10, RGBColor(255, 255, 255), bold=True, align=PP_ALIGN.CENTER)


def add_arrow(slide, x1, y1, x2, y2):
    line = slide.shapes.add_connector(1, x1, y1, x2, y2)
    line.line.color.rgb = MUTED
    line.line.width = Pt(1.5)


def add_line(slide, x, y, w, color):
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, Cm(0.03))
    line.fill.solid()
    line.fill.fore_color.rgb = color
    line.line.color.rgb = color


def add_text(slide, text, x, y, w, h, size, color=INK, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(x, y, w, h)
    frame = box.text_frame
    frame.clear()
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    p = frame.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


if __name__ == "__main__":
    main()
